"""Single source for the state privacy ranking.
Run:  python3 tools/scores.py
Writes public/scores.js (used by the rights page heat map) and public/us-privacy-ranking.xlsx (download).
Review scores and sources each January, or when a state passes a privacy law.
"""
import json, pathlib
ROOT = pathlib.Path(__file__).resolve().parent.parent
REVIEWED = "2026-10-30"
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.datavalidation import DataValidation
F="Arial"; B=Font(name=F,size=10); BB=Font(name=F,size=10,bold=True); H=Font(name=F,size=10,bold=True,color="FFFFFF"); BL=Font(name=F,size=10)
HF=PatternFill("solid",fgColor="3A3A3A"); GF=PatternFill("solid",fgColor="5B6B2E"); th=Side(style="thin",color="C8D1D5"); BD=Border(left=th,right=th,top=th,bottom=th)
TOP=Alignment(vertical="top",wrap_text=True); CTR=Alignment(horizontal="center",vertical="top")

CRIT=[ # key, name, measures, group, weight, scoring guide
("law","Comprehensive law","One law covering most companies and sectors","Consumer",1,"3 broad law in effect · 1 narrow law, or signed but not yet in effect · 0 none"),
("collect","Collection limits","Limits on collecting more data than needed","Consumer",1,"3 strict 'strictly necessary' standard · 2 purpose limits plus rules · 1 'reasonably necessary' standard · 0 none"),
("rights","Individual rights","Rights to see, correct, delete and move your data","Consumer",1,"3 all four · 2 missing correction · 1 narrow law · 0 none"),
("brokers","Data brokers","Brokers registered, covered and removable","Consumer",1,"3 registry plus one-step deletion · 2 registry · 1 covered by law (public-records exemption) · 0 none"),
("sensitive","Sensitive data","Health, biometric, location and children's data","Consumer",1,"3 sale of sensitive data banned · 2 opt-in consent or strong health/biometric law · 1 notice and opt-out · 0 none"),
("regulator","Regulator","Who enforces, and how dedicated they are","Consumer",1,"3 dedicated privacy agency · 2 AG with published privacy enforcement program · 1 AG enforcement · 0 no privacy law to enforce"),
("enforce","Enforcement record","Penalties actually imposed","Consumer",2,"3 repeated, substantial penalties · 2 several actions or large private judgments · 1 at least one public action · 0 none found"),
("remedies","Personal remedies","Complaints, and the right to sue","Consumer",1,"2 private right to sue in some privacy cases · 1 free AG consumer complaints only · 3 broad private right (none yet)"),
("breach","Breach notification","Notice when your data leaks","Consumer",1,"3 law with 30-day deadline · 2 law · 0 none found"),
("const","Constitutional right","Explicit right to privacy in the state constitution","Government",1,"3 explicit and also applies to private companies · 2 explicit, against government · 0 none explicit"),
("police","Police access to data","Warrant rules for location and electronic data, beyond the federal baseline","Government",1,"3 warrant for location (historical and real-time) or electronic data · 2 partial coverage · 1 narrow coverage · 0 federal baseline only"),
]
K=[c[0] for c in CRIT]

# Spanish text for the site (the spreadsheet stays in English)
CRIT_ES={
 "law":("Ley integral","Una ley que cubre la mayoría de las empresas y sectores","3 ley amplia vigente · 1 ley limitada, o aprobada pero aún no vigente · 0 ninguna"),
 "collect":("Límites a la recopilación","Límites a recopilar más datos de los necesarios","3 norma estricta de “estrictamente necesario” · 2 límites de propósito más reglas · 1 norma de “razonablemente necesario” · 0 ninguno"),
 "rights":("Derechos individuales","Derechos a ver, corregir, borrar y trasladar sus datos","3 los cuatro · 2 falta la corrección · 1 ley limitada · 0 ninguno"),
 "brokers":("Corredores de datos","Corredores registrados, cubiertos y con opción de baja","3 registro más eliminación en un solo paso · 2 registro · 1 cubiertos por la ley (excepción de registros públicos) · 0 nada"),
 "sensitive":("Datos sensibles","Datos de salud, biométricos, de ubicación y de menores","3 prohibida la venta de datos sensibles · 2 consentimiento previo o ley fuerte de salud/biometría · 1 aviso y opción de baja · 0 nada"),
 "regulator":("Autoridad","Quién hace cumplir la ley y con qué dedicación","3 agencia de privacidad dedicada · 2 fiscal general con programa público de privacidad · 1 fiscal general · 0 no hay ley de privacidad que aplicar"),
 "enforce":("Historial de sanciones","Sanciones realmente impuestas","3 sanciones repetidas y considerables · 2 varias acciones o grandes fallos privados · 1 al menos una acción pública · 0 ninguna encontrada"),
 "remedies":("Recursos personales","Quejas y derecho a demandar","2 derecho a demandar en algunos casos de privacidad · 1 solo quejas gratuitas ante el fiscal general · 3 derecho amplio a demandar (ninguno todavía)"),
 "breach":("Aviso de filtraciones","Aviso cuando se filtran sus datos","3 ley con plazo de 30 días · 2 ley · 0 ninguna encontrada"),
 "const":("Derecho constitucional","Derecho explícito a la privacidad en la constitución estatal","3 explícito y también se aplica a empresas privadas · 2 explícito, frente al gobierno · 0 ninguno explícito"),
 "police":("Acceso policial a datos","Reglas de orden judicial para datos de ubicación y electrónicos, más allá del mínimo federal","3 orden judicial para ubicación (histórica y en tiempo real) o datos electrónicos · 2 cobertura parcial · 1 cobertura limitada · 0 solo el mínimo federal"),
}
NOTE_ES={
 "Florida Digital Bill of Rights covers only very large companies.":"La Florida Digital Bill of Rights solo cubre a empresas muy grandes.",
 "No right to correct data; weaker collection and sensitive-data rules.":"No hay derecho a corregir datos; reglas más débiles sobre recopilación y datos sensibles.",
 "Strictest collection standard in the US; bans sale of sensitive data.":"La norma de recopilación más estricta del país; prohíbe vender datos sensibles.",
 "Only state with a dedicated privacy agency (CalPrivacy); DROP one-step broker deletion; repeated CCPA fines in 2025; constitutional privacy right applies to companies too.":"Único estado con una agencia de privacidad dedicada (CalPrivacy); eliminación en un solo paso con DROP; multas repetidas bajo la CCPA en 2025; el derecho constitucional a la privacidad también se aplica a las empresas.",
 "Runs a data broker registry.":"Tiene un registro de corredores de datos.",
 "Billion-dollar biometric privacy settlements (Meta 2024, Google 2025) and a TDPSA lawsuit.":"Acuerdos multimillonarios por privacidad biométrica (Meta 2024, Google 2025) y una demanda bajo la TDPSA.",
 "AG publishes annual privacy enforcement reports; CTDPA penalties in 2025.":"El fiscal general publica informes anuales de aplicación de la ley de privacidad; sanciones bajo la CTDPA en 2025.",
 "DOJ publishes annual and quarterly privacy enforcement reports.":"El Departamento de Justicia publica informes anuales y trimestrales de aplicación de la ley de privacidad.",
 "AG sued Snap under the Utah Consumer Privacy Act (2025).":"El fiscal general demandó a Snap bajo la Utah Consumer Privacy Act (2025).",
 "AG sued Roku under the Digital Bill of Rights (2025).":"El fiscal general demandó a Roku bajo la Digital Bill of Rights (2025).",
 "No comprehensive law, but BIPA is the strongest biometric law in the US, with a private right to sue and large judgments.":"No tiene ley integral, pero la BIPA es la ley biométrica más fuerte del país, con derecho a demandar y grandes fallos.",
 "No comprehensive law, but the My Health My Data Act protects health data with a private right to sue.":"No tiene ley integral, pero la My Health My Data Act protege los datos de salud con derecho a demandar.",
 "No comprehensive law; has a consumer health data law.":"No tiene ley integral; tiene una ley de datos de salud del consumidor.",
 "No breach notification law found in the sources reviewed.":"No se encontró una ley de aviso de filtraciones en las fuentes revisadas.",
 "State constitution protects electronic data from warrantless searches.":"La constitución estatal protege los datos electrónicos de registros sin orden judicial.",
}
MONTHS_ES={"January":"enero","February":"febrero","March":"marzo","April":"abril","May":"mayo","June":"junio","July":"julio","August":"agosto","September":"septiembre","October":"octubre","November":"noviembre","December":"diciembre"}
def note_es(note):
    import re as _re
    parts=[x.strip() for x in _re.split(r"(?<=\.)\s+(?=[A-Z])",note) if x.strip()]
    out=[]
    for x in parts:
        m=_re.match(r"Law signed in 2026, takes effect (\w+) (\d+), (\d+); other scores rise then\.",x)
        if m: out.append(f"Ley aprobada en 2026, entra en vigor el {m.group(2)} de {MONTHS_ES[m.group(1)]} de {m.group(3)}; los demás puntajes subirán entonces."); continue
        out.append(NOTE_ES.get(x,x))
    return " ".join(out)

# Base sets
law3={"CA","CO","CT","DE","IN","IA","KY","MD","MN","MT","NE","NH","NJ","OR","RI","TN","TX","UT","VA"}
pending={"AL":"May 1, 2027","LA":"January 1, 2027","OK":"January 1, 2027","VT":"January 1, 2028"}
const2={"AK","AZ","FL","HI","IL","LA","MT","NH","SC","WA","PR"}
police={"CA":3,"ME":3,"MN":3,"MT":3,"NH":3,"UT":3,"MI":3,"MO":3,"CO":2,"FL":2,"IL":2,"IN":2,"MD":2,"MA":2,"NJ":2,"VA":2,"WA":1,"IA":1}
breach30={"CO","FL","WA","ME","NY"}
J=[("AL","Alabama"),("AK","Alaska"),("AZ","Arizona"),("AR","Arkansas"),("CA","California"),("CO","Colorado"),("CT","Connecticut"),("DE","Delaware"),("DC","District of Columbia"),("FL","Florida"),("GA","Georgia"),("HI","Hawaii"),("ID","Idaho"),("IL","Illinois"),("IN","Indiana"),("IA","Iowa"),("KS","Kansas"),("KY","Kentucky"),("LA","Louisiana"),("ME","Maine"),("MD","Maryland"),("MA","Massachusetts"),("MI","Michigan"),("MN","Minnesota"),("MS","Mississippi"),("MO","Missouri"),("MT","Montana"),("NE","Nebraska"),("NV","Nevada"),("NH","New Hampshire"),("NJ","New Jersey"),("NM","New Mexico"),("NY","New York"),("NC","North Carolina"),("ND","North Dakota"),("OH","Ohio"),("OK","Oklahoma"),("OR","Oregon"),("PA","Pennsylvania"),("RI","Rhode Island"),("SC","South Carolina"),("SD","South Dakota"),("TN","Tennessee"),("TX","Texas"),("UT","Utah"),("VT","Vermont"),("VA","Virginia"),("WA","Washington"),("WV","West Virginia"),("WI","Wisconsin"),("WY","Wyoming"),
   ("PR","Puerto Rico"),("GU","Guam"),("VI","U.S. Virgin Islands"),("AS","American Samoa"),("MP","Northern Mariana Islands")]
terr={"PR","GU","VI","AS","MP"}
rows=[]
for c,n in J:
    s={k:0 for k in K}; notes=[]
    if c in law3:
        s.update(law=3,rights=3,collect=1,brokers=1,sensitive=2,regulator=1)
    if c in pending:
        s["law"]=1; notes.append(f"Law signed in 2026, takes effect {pending[c]}; other scores rise then.")
    if c=="FL": s.update(law=1,rights=1,sensitive=1,regulator=1); notes.append("Florida Digital Bill of Rights covers only very large companies.")
    if c in ("UT","IA"): s.update(rights=2,collect=0,sensitive=1); notes.append("No right to correct data; weaker collection and sensitive-data rules.")
    if c=="MD": s.update(collect=3,sensitive=3); notes.append("Strictest collection standard in the US; bans sale of sensitive data.")
    if c=="CA":
        s.update(collect=2,brokers=3,regulator=3,enforce=3,remedies=2,const=3)
        notes.append("Only state with a dedicated privacy agency (CalPrivacy); DROP one-step broker deletion; repeated CCPA fines in 2025; constitutional privacy right applies to companies too.")
    if c in ("VT","TX","OR"): s["brokers"]=2; notes.append("Runs a data broker registry.")
    if c=="TX": s.update(regulator=2,enforce=3); notes.append("Billion-dollar biometric privacy settlements (Meta 2024, Google 2025) and a TDPSA lawsuit.")
    if c=="CT": s.update(regulator=2,enforce=2); notes.append("AG publishes annual privacy enforcement reports; CTDPA penalties in 2025.")
    if c=="OR": s.update(regulator=2,enforce=2); notes.append("DOJ publishes annual and quarterly privacy enforcement reports.")
    if c=="UT": s["enforce"]=1; notes.append("AG sued Snap under the Utah Consumer Privacy Act (2025).")
    if c=="FL": s["enforce"]=1; notes.append("AG sued Roku under the Digital Bill of Rights (2025).")
    if c=="IL": s.update(sensitive=2,enforce=2,remedies=2); notes.append("No comprehensive law, but BIPA is the strongest biometric law in the US, with a private right to sue and large judgments.")
    if c=="WA": s.update(sensitive=2,remedies=2); notes.append("No comprehensive law, but the My Health My Data Act protects health data with a private right to sue.")
    if c=="NV": s["sensitive"]=1; notes.append("No comprehensive law; has a consumer health data law.")
    s["remedies"]=max(s["remedies"],1)
    if c in ("AS","MP"):
        s["breach"]=0; notes.append("No breach notification law found in the sources reviewed.")
    else:
        s["breach"]=3 if c in breach30 else 2
    if c in const2: s["const"]=max(s["const"],2)
    if c in police: s["police"]=police[c]
    if c in ("MI","MO"): notes.append("State constitution protects electronic data from warrantless searches.")
    rows.append((n,"Territory" if c in terr else ("District" if c=="DC" else "State"),s," ".join(notes)))

wb=Workbook()
# ---- Ranking sheet
ws=wb.active; ws.title="Ranking"
hdr=["Jurisdiction","Type"]+[c[1] for c in CRIT]+["Consumer score","Consumer %","Government score","Consumer rank","Notes and judgment calls"]
ws.append(hdr)
nC=len(CRIT); first=3; lastc=first+nC-1
from openpyxl.utils import get_column_letter as L
cols={k:L(first+i) for i,k in enumerate(K)}
cons=[k for k,_,_,g,_,_ in CRIT if g=="Consumer"]; gov=[k for k,_,_,g,_,_ in CRIT if g=="Government"]
cs,cp,gs,rk,nt=[L(lastc+i) for i in range(1,6)]
N=len(rows); end=N+1
for i,(n,t,s,note) in enumerate(rows):
    r=i+2
    ws.append([n,t]+[s[k] for k in K])
    terms=[f"{cols[k]}{r}*Criteria!$E${2+K.index(k)}" for k in cons]
    ws[f"{cs}{r}"]="="+"+".join(terms)
    ws[f"{cp}{r}"]=f"={cs}{r}/Criteria!$F${len(CRIT)+3}"
    ws[f"{gs}{r}"]="="+"+".join(f"{cols[k]}{r}" for k in gov)
    ws[f"{rk}{r}"]=f"=RANK({cs}{r},${cs}$2:${cs}${end})"
    ws[f"{nt}{r}"]=note
for row in ws.iter_rows(min_row=1,max_row=end):
    for c in row: c.font=B; c.alignment=TOP; c.border=BD
for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
for k in gov: ws[f"{cols[k]}1"].fill=GF
ws[f"{gs}1"].fill=GF
for r in range(2,end+1):
    for k in K: ws[f"{cols[k]}{r}"].font=BL; ws[f"{cols[k]}{r}"].alignment=CTR
    for col in (cs,cp,gs,rk): ws[f"{col}{r}"].font=BB; ws[f"{col}{r}"].alignment=CTR
    ws[f"{cp}{r}"].number_format="0%"
HEAT=dict(start_type="num",start_value=0,start_color="F4F1E8",mid_type="num",mid_value=1.5,mid_color="E9C46A",end_type="num",end_value=3,end_color="2E9E5B")
ws.conditional_formatting.add(f"{cols[K[0]]}2:{cols[K[-1]]}{end}",ColorScaleRule(**HEAT))
ws.conditional_formatting.add(f"{cp}2:{cp}{end}",ColorScaleRule(start_type="num",start_value=0,start_color="F4F1E8",mid_type="num",mid_value=0.4,mid_color="E9C46A",end_type="num",end_value=0.8,end_color="2E9E5B"))
ws.conditional_formatting.add(f"{gs}2:{gs}{end}",ColorScaleRule(start_type="num",start_value=0,start_color="F4F1E8",mid_type="num",mid_value=3,mid_color="E9C46A",end_type="num",end_value=6,end_color="2E9E5B"))
dv=DataValidation(type="whole",operator="between",formula1="0",formula2="3",showErrorMessage=True,error="Scores are whole numbers from 0 to 3."); ws.add_data_validation(dv); dv.add(f"{cols[K[0]]}2:{cols[K[-1]]}{end}")
ws.column_dimensions["A"].width=22; ws.column_dimensions["B"].width=9
for k in K: ws.column_dimensions[cols[k]].width=11
for col in (cs,cp,gs,rk): ws.column_dimensions[col].width=11
ws.column_dimensions[nt].width=70
ws.row_dimensions[1].height=44
ws.freeze_panes="C2"; ws.auto_filter.ref=f"A1:{nt}{end}"
ws[f"{cols['police']}1"].comment=Comment("Lower confidence: based mainly on a 2015 ACLU survey of state location-privacy laws; some states may have added protections since. The federal Carpenter decision (2018) already requires a warrant for historical cell-site records everywhere.","Opt Out Kit")
ws[f"{cols['enforce']}1"].comment=Comment("Weighted double. Based on public actions reported through 2025; states with new laws may simply not have acted yet.","Opt Out Kit")
ws[f"{cols['const']}1"].comment=Comment("Most explicit state constitutional rights limit government, not companies. California's also applies to private parties.","Opt Out Kit")

# ---- Criteria sheet
c2=wb.create_sheet("Criteria")
c2.append(["#","Criterion","What it measures","Score type","Weight","Max points","How it's scored"])
for i,(k,n,m,g,w,guide) in enumerate(CRIT,1):
    c2.append([i,n,m,g,w,f"=3*E{i+1}",guide])
e=len(CRIT)+1
c2.append([])
c2.append(["","Consumer protection max","","",f'=SUMIF(D2:D{e},"Consumer",E2:E{e})',f'=SUMIF(D2:D{e},"Consumer",F2:F{e})'])
c2.append(["","Government max","","",f'=SUMIF(D2:D{e},"Government",E2:E{e})',f'=SUMIF(D2:D{e},"Government",F2:F{e})'])
for row in c2.iter_rows(min_row=1,max_row=c2.max_row):
    for c in row: c.font=B; c.alignment=TOP
for c in c2[1]: c.font=H; c.fill=HF
for r in range(2,e+1):
    for c in c2[r]: c.border=BD
    c2.cell(r,5).font=BL
for r in (e+2,e+3):
    for c in c2[r]: c.font=BB
for col,wd in zip("ABCDEFG",[4,22,40,12,8,10,70]): c2.column_dimensions[col].width=wd
c2.freeze_panes="A2"

# ---- About sheet
a=wb.create_sheet("About")
lines=[("U.S. state and territory privacy ranking",True),
("Scores reflect protections in effect as of "+REVIEWED+", measured against the EU's GDPR as the benchmark. Laws signed but not yet in effect are noted and score 1 on 'Comprehensive law'.",False),
("Consumer score covers protection from companies (max 30, enforcement weighted double). Government score covers constitutional privacy rights and police access to data (max 6) and is kept separate so it doesn't mask the consumer score.",False),
("Score cells are shaded as a heat map, pale cream (0) through gold to green (3). They can be edited (whole numbers 0-3); totals, percentages and ranks recalculate. Change weights on the Criteria tab.",False),
("Many scores are judgment calls. The Notes column explains each state's notable features. Where no specific evidence was found, a state gets the default score for its category: for example, every state with a broad law in effect starts at collection 1, rights 3, brokers 1, sensitive 2, regulator 1.",False),
("Known limits: police-access scores rely mainly on a 2015 survey and may miss newer laws; enforcement reflects public actions through 2025; breach-law 30-day deadlines are confirmed only for the states marked; American Samoa and the Northern Mariana Islands had no breach law in the sources reviewed, which may reflect a gap in those sources.",False),
("Sources",True),
("State comprehensive laws and dates: ClauseAudit state privacy law tracker; Morgan Lewis, 'New US State Consumer Privacy Laws' (July 2026); MultiState (Oct. 28, 2026).",False),
("Enforcement: Future of Privacy Forum, 'U.S. Privacy Enforcement in 2025' (Feb. 2026); MultiState, 'How state AGs are policing consumer data practices' (Sept. 2026).",False),
("Police access: ACLU, 'Status of Location Privacy Legislation in the States' (2015); EPIC, State Location Privacy Policy.",False),
("Breach notification: DLA Piper, Data Protection Laws of the World (U.S. breach notification).",False),
("General information, not legal advice. Laws change every legislative session.",False)]
for i,(t,b) in enumerate(lines,1):
    c=a.cell(i,1,t); c.font=Font(name=F,size=12 if i==1 else 10,bold=b); c.alignment=TOP
a.column_dimensions["A"].width=110
wb.move_sheet("About",offset=-2)
wb.active=1
wb.save(ROOT / "public/us-privacy-ranking.xlsx")

# ---- scores.js for the site
codes=[c for c,_ in J]
out={"reviewed":REVIEWED,
     "criteria":[{"k":k,"name":n,"what":m,"group":g,"weight":w,"guide":guide,"name_es":CRIT_ES[k][0],"what_es":CRIT_ES[k][1],"guide_es":CRIT_ES[k][2]} for k,n,m,g,w,guide in CRIT],
     "scores":{c:{"s":[rows[i][2][k] for k in K],"note":rows[i][3],"note_es":note_es(rows[i][3])} for i,c in enumerate(codes)}}
(ROOT / "public/scores.js").write_text("/* Generated by tools/scores.py. Do not edit by hand. */\nwindow.OOK_SCORES = "+json.dumps(out,ensure_ascii=False,separators=(",",":"))+";\n",encoding="utf-8")
print(len(rows),"jurisdictions scored; wrote public/scores.js and public/us-privacy-ranking.xlsx")
